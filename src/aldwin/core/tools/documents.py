"""`read_document`: convert PDF/Office/CSV/HTML to markdown for the model."""

from __future__ import annotations

import asyncio
import csv
import io
import re
from pathlib import Path
from typing import Any, Final

import markdownify
import pypdfium2 as pdfium
from agentling import Tool
from agentling.tools import ToolCallError
from bs4 import BeautifulSoup

from aldwin.core.policy import PolicySession
from aldwin.core.tools import ToolOutput
from aldwin.core.tools.files import format_size
from aldwin.core.workspace import Workspace, WorkspaceError

DOC_TIMEOUT_SECONDS: Final[float] = 60.0
MAX_DOCUMENT_BYTES: Final[int] = 50 * 1024 * 1024
MAX_DOCUMENT_CHARS: Final[int] = 30_000
MAX_TABLE_ROWS: Final[int] = 200
SCANNED_CHAR_THRESHOLD: Final[int] = 20
SUPPORTED_EXTENSIONS: Final[frozenset[str]] = frozenset(
    {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".tsv", ".html", ".htm"}
)
DOCUMENT_TOOL_NAMES: Final[frozenset[str]] = frozenset({"read_document"})

_OBJECT_SCHEMA: Final[dict[str, Any]] = {
    "type": "object",
    "additionalProperties": False,
}
_SUPPORTED_LIST: Final[str] = "pdf, docx, pptx, xlsx, csv, tsv, html"
_OFFICE_EXTRA_HINT: Final[str] = (
    'needs the office extra: install with `uv tool install "aldwin[office]"` '
    '(or `pip install "aldwin[office]"`)'
)


def running_summary(
    name: str,
    arguments: dict[str, Any] | None,
    workspace: Workspace,
) -> str:
    """Pending row for `read_document`, or `\"\"` for other tools."""
    if name not in DOCUMENT_TOOL_NAMES:
        return ""
    args = arguments or {}
    raw = args.get("path", "")
    shown = raw if isinstance(raw, str) else str(raw)
    try:
        display = workspace.resolve_safe(shown).display
    except WorkspaceError:
        display = shown or "document"
    return f"Read {display}"


def row_summary(running: str, summary: str, *, ok: bool) -> str:
    """Rewrite a denial row for document tools."""
    if ok or not running:
        return summary
    lower = summary.lower()
    if summary == "declined" or "denied by the user" in lower:
        return f"{running} · declined"
    if summary == "not run: no approver" or "no one was available to approve" in lower:
        return f"{running} · not run: no one to approve"
    if (summary.startswith("blocked:") and "secret" in lower) or (
        "blocked because it may hold secrets" in lower
    ):
        return f"{running} · blocked: may contain secrets"
    return summary


def parse_pages_spec(spec: str, page_count: int) -> tuple[list[int], list[str]]:
    """Parse a 1-based `pages` string into page indices and out-of-range notes."""
    selected: list[int] = []
    notes: list[str] = []
    seen: set[int] = set()
    for part in spec.split(","):
        token = part.strip()
        if not token:
            continue
        if "-" in token:
            left, _, right = token.partition("-")
            try:
                start = int(left.strip())
                end = int(right.strip())
            except ValueError as exc:
                raise ToolCallError(f"invalid pages: {spec!r}") from exc
            if start > end:
                start, end = end, start
            for page in range(start, end + 1):
                if page < 1 or page > page_count:
                    notes.append(f"page {page} does not exist ({page_count} pages)")
                elif page not in seen:
                    selected.append(page)
                    seen.add(page)
        else:
            try:
                page = int(token)
            except ValueError as exc:
                raise ToolCallError(f"invalid pages: {spec!r}") from exc
            if page < 1 or page > page_count:
                notes.append(f"page {page} does not exist ({page_count} pages)")
            elif page not in seen:
                selected.append(page)
                seen.add(page)
    return selected, notes


def _word_count(text: str) -> int:
    return len(re.findall(r"\S+", text))


def _page_image_count(page: Any) -> int:
    try:
        objects = list(page.get_objects())
    except Exception:
        return 0
    count = 0
    for obj in objects:
        name = type(obj).__name__.lower()
        if "image" in name:
            count += 1
    return count


def _extract_pdf_pages(
    path: Path,
    *,
    pages_spec: str | None,
    cancel: asyncio.Event | None = None,
) -> tuple[str, str]:
    """Return `(markdown_body, summary_suffix)` for a PDF."""
    try:
        pdf = pdfium.PdfDocument(path)
    except pdfium.PdfiumError as exc:
        code = getattr(exc, "err_code", None)
        if code == 4:
            raise ToolCallError(
                f"encrypted: {path.name}; open it and save an unlocked copy"
            ) from exc
        if code == 3:
            raise ToolCallError(f"not a valid PDF: {path.name}") from exc
        raise ToolCallError(f"failed to open PDF: {exc}") from exc

    page_count = len(pdf)
    if pages_spec:
        selected, notes = parse_pages_spec(pages_spec, page_count)
        if not selected and notes:
            body = "\n".join(notes)
            return body, "pages none · 0 words"
    else:
        selected = list(range(1, page_count + 1))
        notes = []

    chunks: list[str] = []
    total_words = 0
    shown_pages: list[int] = []
    truncated_from: int | None = None
    for index, page_no in enumerate(selected):
        if cancel is not None and cancel.is_set():
            raise asyncio.CancelledError()
        page = pdf[page_no - 1]
        textpage = page.get_textpage()
        raw = textpage.get_text_bounded() or ""
        non_ws = len(re.sub(r"\s+", "", raw))
        images = _page_image_count(page)
        page_chunks = [f"## Page {page_no}"]
        if non_ws < SCANNED_CHAR_THRESHOLD and images >= 1:
            page_chunks.append("(no text layer: scanned page; OCR arrives in 0.0.3)")
            page_words = 0
        elif raw.strip():
            page_chunks.append(raw.strip())
            page_words = _word_count(raw)
        else:
            page_chunks.append("")
            page_words = 0
        page_chunks.append("")
        candidate = "\n".join(chunks + page_chunks)
        if shown_pages and len(candidate) > MAX_DOCUMENT_CHARS:
            truncated_from = page_no
            remaining = selected[index:]
            if remaining:
                first_omit, last_omit = remaining[0], remaining[-1]
                chunks.append(
                    f"… pages {first_omit}–{last_omit} not shown; "
                    f'call again with pages="{first_omit}-{last_omit}"'
                )
            break
        chunks.extend(page_chunks)
        shown_pages.append(page_no)
        total_words += page_words

    if truncated_from is None:
        for note in notes:
            chunks.append(note)

    body = "\n".join(chunks).rstrip() + "\n"
    if shown_pages:
        first, last = shown_pages[0], shown_pages[-1]
        suffix = f"pages {first}–{last} · {total_words:,} words"
    else:
        suffix = "pages none · 0 words"
    return body, suffix


def _sheet_to_markdown(rows: list[list[Any]], *, total_rows: int) -> str:
    if not rows:
        return ""
    header = [str(c) if c is not None else "" for c in rows[0]]
    if not any(header):
        header = [f"col{i + 1}" for i in range(len(header) or 1)]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    data_rows = rows[1:]
    shown = data_rows[:MAX_TABLE_ROWS]
    for row in shown:
        cells: list[str] = []
        for cell in row:
            if cell is None:
                cells.append("")
            elif isinstance(cell, float) and cell.is_integer():
                cells.append(str(int(cell)))
            else:
                cells.append(str(cell))
        while len(cells) < len(header):
            cells.append("")
        lines.append("| " + " | ".join(cells[: len(header)]) + " |")
    remaining = max(0, total_rows - 1 - len(shown))  # minus header
    if remaining > 0:
        lines.append(f"… {remaining:,} more rows")
    return "\n".join(lines)


def _read_xlsx(path: Path, sheet: str | None) -> tuple[str, str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ToolCallError(f"reading .xlsx {_OFFICE_EXTRA_HINT}") from exc

    wb_form = load_workbook(path, read_only=True, data_only=False)
    try:
        names = list(wb_form.sheetnames)
        if not names:
            raise ToolCallError(f"no sheets in {path.name}")

        if sheet is None or sheet.strip() == "":
            chosen = names[0]
        elif sheet.isdigit():
            index = int(sheet)
            if index < 1 or index > len(names):
                raise ToolCallError(
                    f"sheet {sheet!r} does not exist; Sheets: {', '.join(names)}"
                )
            chosen = names[index - 1]
        elif sheet in names:
            chosen = sheet
        else:
            raise ToolCallError(
                f"sheet {sheet!r} does not exist; Sheets: {', '.join(names)}"
            )

        form_rows = [list(row) for row in wb_form[chosen].iter_rows(values_only=True)]
    finally:
        wb_form.close()

    wb_data = load_workbook(path, read_only=True, data_only=True)
    try:
        data_rows = [list(row) for row in wb_data[chosen].iter_rows(values_only=True)]
    finally:
        wb_data.close()

    # Prefer cached values; mark missing formula caches.
    merged: list[list[Any]] = []
    max_len = max(len(data_rows), len(form_rows))
    for i in range(max_len):
        drow = data_rows[i] if i < len(data_rows) else []
        frow = form_rows[i] if i < len(form_rows) else []
        width = max(len(drow), len(frow))
        out_row: list[Any] = []
        for j in range(width):
            dval = drow[j] if j < len(drow) else None
            fval = frow[j] if j < len(frow) else None
            if dval is None and isinstance(fval, str) and fval.startswith("="):
                out_row.append("(formula, not calculated)")
            else:
                out_row.append(dval)
        merged.append(out_row)

    total_data = max(0, len(merged) - 1)
    table = _sheet_to_markdown(merged, total_rows=len(merged))
    header = f"Sheets: {', '.join(names)}"
    body = f"{header}\n\n{table}\n" if table else f"{header}\n"
    shown = min(MAX_TABLE_ROWS, total_data)
    suffix = f"sheet {chosen} · {shown} of {total_data:,} rows"
    return body, suffix


def _read_csv(path: Path, *, delimiter: str | None = None) -> tuple[str, str]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    sample = text[:4096]
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = "\t" if path.suffix.lower() == ".tsv" else ","
    else:
        delim = delimiter
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = [list(row) for row in reader]
    total_data = max(0, len(rows) - 1)
    table = _sheet_to_markdown(rows, total_rows=len(rows))
    shown = min(MAX_TABLE_ROWS, total_data)
    return table + "\n", f"{shown} of {total_data} rows"


def _read_docx(path: Path) -> tuple[str, str]:
    try:
        import mammoth
    except ImportError as exc:
        raise ToolCallError(f"reading .docx {_OFFICE_EXTRA_HINT}") from exc

    with path.open("rb") as handle:
        result = mammoth.convert_to_markdown(handle)
    text = (result.value or "").strip() + "\n"
    return text, f"{_word_count(text):,} words"


def _read_pptx(path: Path) -> tuple[str, str]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ToolCallError(f"reading .pptx {_OFFICE_EXTRA_HINT}") from exc

    prs = Presentation(str(path))
    chunks: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        chunks.append(f"## Slide {index}")
        texts: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if line:
                    texts.append(line)
        if texts:
            chunks.extend(texts)
        chunks.append("")
    body = "\n".join(chunks).rstrip() + "\n"
    return body, f"{len(prs.slides)} slides · {_word_count(body):,} words"


def _read_html(path: Path) -> tuple[str, str]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "aside"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    md = markdownify.markdownify(str(main), heading_style="ATX", strip=["img"])
    text = (md or "").strip() + "\n"
    return text, f"{_word_count(text):,} words"


def convert_document(
    path: Path,
    *,
    pages: str | None = None,
    sheet: str | None = None,
    cancel: asyncio.Event | None = None,
) -> tuple[str, str]:
    """Convert `path` to markdown and a one-line CLI summary suffix."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ToolCallError(
            f"unsupported type {ext or path.name}; supported: {_SUPPORTED_LIST}"
        )
    size = path.stat().st_size
    if size > MAX_DOCUMENT_BYTES:
        raise ToolCallError(f"too large: {format_size(size)} (limit 50 MB)")
    if ext == ".pdf":
        return _extract_pdf_pages(path, pages_spec=pages, cancel=cancel)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".pptx":
        return _read_pptx(path)
    if ext == ".xlsx":
        return _read_xlsx(path, sheet)
    if ext in {".csv", ".tsv"}:
        return _read_csv(path, delimiter="\t" if ext == ".tsv" else None)
    if ext in {".html", ".htm"}:
        return _read_html(path)
    raise ToolCallError(f"unsupported type {ext}; supported: {_SUPPORTED_LIST}")


class ReadDocumentTool(Tool):
    """Read a document file as markdown text."""

    name = "read_document"
    description = (
        "Read a PDF, Word, PowerPoint, Excel, CSV/TSV, or HTML file as markdown. "
        "Use pages= for PDF page ranges and sheet= for Excel sheet name or "
        "1-based index."
    )
    parallel_safe = True
    timeout = None
    max_output_chars = 30_000
    parameters: dict[str, Any] = {
        **_OBJECT_SCHEMA,
        "properties": {
            "path": {
                "type": "string",
                "description": "Document path, relative to the workspace.",
            },
            "pages": {
                "type": "string",
                "description": 'PDF pages to read, e.g. "1-5,9" (1-based).',
            },
            "sheet": {
                "type": "string",
                "description": "Excel sheet name or 1-based index.",
            },
        },
        "required": ["path"],
    }

    def __init__(self, workspace: Workspace, policy_session: PolicySession) -> None:
        self._workspace = workspace
        self._session = policy_session

    async def forward(  # type: ignore[override]
        self,
        path: str,
        pages: str | None = None,
        sheet: str | None = None,
    ) -> ToolOutput:
        try:
            resolved = self._workspace.resolve_safe(
                path,
                granted_folders=self._session.granted_folders,
            )
        except WorkspaceError as exc:
            raise ToolCallError(str(exc)) from exc

        target = resolved.path
        display = resolved.display
        if not target.exists():
            raise ToolCallError(f"not found: {display}")
        if target.is_dir():
            raise ToolCallError(f"is a directory: {display}")

        cancel = asyncio.Event()

        def _run() -> tuple[str, str]:
            return convert_document(
                target,
                pages=pages,
                sheet=sheet,
                cancel=cancel,
            )

        try:
            async with asyncio.timeout(DOC_TIMEOUT_SECONDS):
                body, suffix = await asyncio.to_thread(_run)
        except TimeoutError as exc:
            cancel.set()
            raise ToolCallError(f"timed out after {int(DOC_TIMEOUT_SECONDS)}s") from exc
        except asyncio.CancelledError:
            cancel.set()
            raise

        summary = f"Read {display} · {suffix}"
        return ToolOutput(body, summary=summary, ok=True)


def build_document_tools(
    workspace: Workspace,
    policy_session: PolicySession,
) -> list[Tool]:
    """Return the `read_document` tool."""
    return [ReadDocumentTool(workspace, policy_session)]


__all__ = [
    "DOCUMENT_TOOL_NAMES",
    "ReadDocumentTool",
    "build_document_tools",
    "convert_document",
    "parse_pages_spec",
    "row_summary",
    "running_summary",
]
